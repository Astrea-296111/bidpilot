"""Generate selectable-text PDF and DOCX demo fixtures, without platform fonts."""

from pathlib import Path

import fitz
from docx import Document


def main():
    root = Path(__file__).resolve().parents[1]
    source = root / "data/rfps/01-bank-data-platform.md"
    text = source.read_text(encoding="utf-8")
    pdf = fitz.open()
    page = pdf.new_page()
    y = 45
    for line in text.splitlines():
        if y > 760:
            page = pdf.new_page()
            y = 45
        page.insert_text((35, y), line, fontname="china-s", fontsize=10)
        y += 20
    pdf.save(source.with_suffix(".pdf"))
    pdf.close()
    doc = Document()
    for line in text.splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:], 0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 1)
        else:
            doc.add_paragraph(line)
    doc.save(source.with_suffix(".docx"))
    print("Generated PDF and DOCX:", source.stem)


if __name__ == "__main__":
    main()
